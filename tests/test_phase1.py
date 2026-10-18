#!/usr/bin/env python3
"""Phase-1 end-to-end test.

1. build the sample catalogue workbook from the JSON fixture, run checker (must be error-free)
2. (seeded-error detection: see tests/test_model.py)
3. init + build sample project, enter sample CI data, rebuild (preservation + manual override)
4. recalculate in LibreOffice, zero formula errors
5. compare every CI / CR / SR / SP status and presumption flag with the Python reference roll-up
6. change a REGISTER label in the catalogue, rebuild, and confirm formulas follow the new label
usage: python tests/test_phase1.py OUTDIR RECALC_SCRIPT
"""
import json, os, shutil, subprocess, sys
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
from openpyxl import load_workbook
import make_catalogue, check_catalogue, build_project
from ctm import vocab
from ctm.store import load_catalogue, load_project_seed, to_tables
from ctm.schema import CATALOGUE, CI_COLS
from ctm.common import read_table
from ctm.register import read_register
from ctm import rollup

OUT, RECALC = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
FIX = os.path.join(HERE, 'fixtures')
SEED = load_project_seed(os.path.join(FIX, 'sample_project.json'))
_lists = dict(verdict='VERDICT', finding_severity='SEV', applicability='APPL')
CI_ENTRIES = {cid: {k: vocab.label_of(_lists[k], v) if k in _lists else v for k, v in d.items()}   # keys -> labels as typed
              for cid, d in json.load(open(os.path.join(FIX, 'sample_ci_entries.json'), encoding='utf-8')).items()}
ok = True


def say(flag, msg):
    global ok
    ok &= flag
    print(('PASS ' if flag else 'FAIL ') + msg)


def recalc(path):
    res = json.loads(subprocess.run([sys.executable, RECALC, path, '120'], capture_output=True, text=True).stdout)
    return res


cat = os.path.join(OUT, 'Catalogue_SAMPLE.xlsx')
make_catalogue.build(cat, to_tables(load_catalogue(os.path.join(FIX, 'sample', 'catalogue'), [])))
E, W, st = check_catalogue.check(os.path.join(FIX, 'sample', 'catalogue'), [])
say(not E, f'catalogue check: {len(E)} errors, {len(W)} warnings {st}')

# 3. project
prj = os.path.join(OUT, 'Project_P26001_SAMPLE.xlsx')
for f in (prj, prj + '.bak'):
    if os.path.exists(f): os.remove(f)
build_project.init_project(cat, prj, SEED)
build_project.build_project(cat, prj)
wb = load_workbook(prj)
ws = wb['CI']
hdr = [c.value for c in ws[1]]
col = {c.key: hdr.index(c.label) + 1 for c in CI_COLS}
for r in range(2, ws.max_row + 1):
    cid = ws.cell(r, col['ci_id']).value
    for k, v in CI_ENTRIES.get(cid, {}).items():
        ws.cell(r, col[k], value=v)
wb.save(prj)
info = build_project.build_project(cat, prj)          # rebuild must keep entries
say(info['orphaned'] == 0 and info['ci'] == 13, f'rebuild: {info}')

# 4. recalc
res = recalc(prj)
say(res.get('status') == 'success', f'recalc project: {res.get("status")} errors={res.get("total_errors")} formulas={res.get("total_formulas")}')
res_c = recalc(cat)
say(res_c.get('status') == 'success', f'recalc catalogue: {res_c.get("status")} errors={res_c.get("total_errors")}')

# 5. compare with reference
v = load_workbook(prj, data_only=True)
reg = read_register(v)
key_of = {lst: {e['label']: e['key'] for e in reg[lst]} for lst in reg}
ci_rows = read_table(v, 'CI', CI_COLS)
kept = all(ci_rows[[c['ci_id'] for c in ci_rows].index(cid)][k] == val for cid, d in CI_ENTRIES.items() for k, val in d.items())
say(kept, 'CI entries preserved across rebuild (incl. manual applicability overrides)')
C = {s: read_table(v, f'cat_{s}', cols) for s, cols in CATALOGUE.items()}
basis = [b.sd_id for b in SEED.basis]
sps = {r['sp_id']: r for r in C['SP'] if r['sd_id'] in basis}
srs = {}
for r in C['SR']:
    refs = build_project.split_refs(r['sp_refs'])
    if refs[0] in sps: srs[r['sr_id']] = dict(r, refs=refs)
crs_ids = [row[0].value for row in v['V_CR'].iter_rows(min_row=2) if row[0].value]
crs = {r['cr_id']: r for r in C['CR'] if r['cr_id'] in crs_ids}
links = [l for l in C['LINK_SR_CR'] if l['sr_id'] in srs]
ref = rollup.compute(reg, ci_rows, crs, links, srs, sps, C['CORR'], {r['sd_id']: r['lifecycle_status'] for r in C['SD']},
                     {r['cr_id']: r for r in C['CR']})

def sheet_map(name, idcol, stcol):
    ws = v[name]; h = [c.value for c in ws[1]]
    return {row[h.index(idcol)].value: row[h.index(stcol)].value for row in ws.iter_rows(min_row=2) if row[0].value}

mism = []
for c in ci_rows:
    exp = ref['ci'][c['ci_id']]
    got = c['status']
    exp_label = '' if not exp else [e['label'] for e in reg['VERDICT' if exp[0] == 'V' else 'APPL'] if e['key'] == exp[1]][0]
    if got != exp_label: mism.append(('CI', c['ci_id'], got, exp_label))
for name, idc, stc, refk in [('V_CR', 'CR ID', 'Final status', 'cr'), ('V_SR', 'SR ID', 'SR status', 'sr'),
                             ('V_SP', 'SP ID', 'Final status', 'sp')]:
    got = sheet_map(name, idc, stc)
    for k, exp in ref[refk].items():
        g = key_of['ROLLUP'].get(got.get(k), got.get(k))
        if g != exp: mism.append((name, k, got.get(k), exp))
pres = sheet_map('V_PRESUMPTION', 'Corr ID', 'Presumption claimable?')
for k, exp in ref['pres'].items():
    if pres.get(k) != exp: mism.append(('PRES', k, pres.get(k), exp))
n_checked = len(ci_rows) + len(ref['cr']) + len(ref['sr']) + len(ref['sp']) + len(ref['pres'])
say(not mism, f'Excel roll-up == reference for {n_checked} statuses' + (f' MISMATCHES: {mism}' if mism else ''))

# expected scenario outcomes (spec semantics, written out by hand)
exp_cases = {('cr', 'CR.ROB.001'): 'PASS', ('cr', 'CR.PERF.004'): 'OPEN', ('cr', 'CR.PHYS.001'): 'NOTASS',
             ('cr', 'CR.PMM.001'): 'PASS', ('cr', 'CR.HO.001'): 'NA', ('sr', 'SR.AIA.0005'): 'GAP',
             ('sp', 'MR:III-1.5.9'): 'OOS', ('sp', 'MR:III-A(e)'): 'NOREQ', ('sp', 'MR:III-1.2.1'): 'OPEN',
             ('sp', 'MR:III-1.2.6'): 'OPEN', ('pres', 'C0001'): 'no', ('pres', 'C0002'): 'yes',
             ('cr', 'CR.LOG.001'): 'OPEN'}
bad_cases = [(k, ref[k[0]][k[1]], e) for k, e in exp_cases.items() if ref[k[0]][k[1]] != e]
say(not bad_cases, f'{len(exp_cases)} hand-written scenario expectations' + (f' {bad_cases}' if bad_cases else ''))
for (kind, k), e in exp_cases.items():
    print(f'      {kind:4} {k:18} -> {ref[kind][k]}')

# 6. register single source of truth: rename a label in the catalogue, rebuild, recalc
cat2 = os.path.join(OUT, '_cat_relabel.xlsx'); prj2 = os.path.join(OUT, '_prj_relabel.xlsx')
shutil.copyfile(cat, cat2); shutil.copyfile(prj, prj2)
wbc = load_workbook(cat2)
for row in wbc['REGISTER'].iter_rows(min_row=2):
    if row[0].value == 'ROLLUP' and row[1].value == 'OPEN':
        row[2].value = 'In progress'
wbc.save(cat2)
build_project.build_project(cat2, prj2)
recalc(prj2)
v2 = load_workbook(prj2, data_only=True)
got = sheet_map.__wrapped__ if hasattr(sheet_map, '__wrapped__') else None
ws2 = v2['V_CR']; h = [c.value for c in ws2[1]]
perf = [r[h.index('Final status')].value for r in ws2.iter_rows(min_row=2) if r[0].value == 'CR.PERF.004'][0]
say(perf == 'In progress', f'relabelled ROLLUP_OPEN propagates to formulas (CR.PERF.004 -> {perf!r})')
for f in (cat2, prj2, prj2 + '.bak'):
    if os.path.exists(f): os.remove(f)
if os.path.exists(prj + '.bak'): os.remove(prj + '.bak')
print('\nALL PASS' if ok else '\nSOME CHECKS FAILED')
sys.exit(0 if ok else 1)
