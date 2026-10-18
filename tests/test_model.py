"""Data model and JSON store tests (no LibreOffice needed).  usage: pytest tests/test_model.py"""
import json, os, shutil, sys, typing

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, 'tools')]

from ctm import checks, model, rules, schemas, vocab                                # noqa: E402
from ctm.schema import CATALOGUE, PROJECT_INPUT                                      # noqa: E402
from ctm.store import (CatalogueError, catalogue_files, default_packs, dumps_model, load_catalogue,  # noqa: E402
                       load_project_seed, read_model, resolve_packs, to_tables)
from ctm.vocab import VocabRef                                                       # noqa: E402
from backbone_metrics import metrics                                                 # noqa: E402

CAT = os.path.join(ROOT, 'catalogue')                     # master list; its public packs are in ../packs
FIX = os.path.join(ROOT, 'tests', 'fixtures')
SAMPLE = os.path.join(FIX, 'sample', 'catalogue')         # sample master list; public packs in sample/packs
EXTRA = os.path.join(FIX, 'extra_packs')                  # directory of extra (private) packs: DEMO
DEMO = os.path.join(EXTRA, 'DEMO')                        # synthetic licensed pack (spec § 8.3)
SEEDS = [os.path.join(FIX, 'sample_project.json')]


@pytest.fixture(scope='module')
def cat():
    return load_catalogue(CAT, [])                       # public packs only, whatever ctm.local.json / CTM_PACKS say


def edit(root, rel, fn):
    p = os.path.join(root, rel)
    d = json.load(open(p, encoding='utf-8'))
    fn(d)
    json.dump(d, open(p, 'w', encoding='utf-8'), ensure_ascii=False)


# ------------------------------------------------------------------------------------------------ public catalogue
# Master list + public packs (MR, AIA, MR-AIDA). The full catalogue with the private packs TS22440-1 and TR5469 is
# tested in the private repository (tests/test_full_catalogue.py).
def test_catalogue_check_clean(cat):
    E, W, stats = checks.check(cat)
    assert (E, W) == ([], [])
    assert stats == dict(PACKS=3, SD=3, SP=1424, SR=335, CR=308, UNREACHED=121, LINK=384, COND=0, CORR=0)
    assert cat.missing_packs() == ['TS22440-1', 'TR5469']


def test_checkpoint_metrics(cat):
    m = metrics(cat)
    assert m['sr'] == {'MR': 221, 'AIA': 114}
    assert m['stats']['MR'] == dict(listed=1244, atomised=250, oos=490, structural=504)
    assert m['stats']['AIA'] == dict(listed=180, atomised=120, oos=0, structural=60)
    assert m['links_by_type'] == {'Covers': 302, 'Covers partly': 75, 'Supports – method': 0, 'Supports – process': 7,
                                  'Context': 0}
    assert (m['cr_total'], m['cr_parents'], m['cr_leaves']) == (308, 23, 285)
    assert m['cr_reached_by'] == {'AIA+MR': 45, 'MR': 73, 'AIA': 66}
    assert m['reuse'] == {'AIA': dict(links=126, onto_first=53)}
    assert len(m['cr_unreached']) == 110 and m['packs_missing'] == ['TS22440-1', 'TR5469']   # leaves of the private packs


def test_public_tree_has_no_licensed_content(cat):
    assert list(cat.packs) == ['MR', 'AIA', 'MR-AIDA'] and set(cat.sds) == {'MR', 'AIA', 'MR-AIDA'}
    assert all(s.licence_class == 'PUBLIC' for s in cat.sds.values())
    for d in os.listdir(os.path.join(ROOT, 'packs')):                 # every pack in this repository is public
        if os.path.isdir(os.path.join(ROOT, 'packs', d)):
            assert any(k.pack_id == d and k.access == 'public' for k in cat.manifest.known_packs), d
    rated = ('AI.usage_level', 'AI.tech_class', 'INT.AISC')            # source-specific: on the links (D0046)
    assert not any(t in c.applicability_rule for c in cat.crs.values() for t in rated)


def test_amendments_applied(cat):
    sps = cat.sps
    assert 'AIA:10(5)' not in sps and 'AIA:10(5)(a)' not in sps
    assert all(p.verbatim for p in sps.values())
    art4a = [k for k in sps if k.startswith('AIA:4a')]
    assert art4a == (['AIA:4a', 'AIA:4a(1)'] + [f'AIA:4a(1)({x})' for x in 'abcdef'] +
                     ['AIA:4a(2)', 'AIA:4a(2)(a)', 'AIA:4a(2)(b)', 'AIA:4a(2)-2'])
    assert sps['AIA:4a(1)(f)'].parent_sp == 'AIA:4a(1)' and sps['AIA:4a(2)-2'].amended_by
    assert list(sps).index('AIA:4a') == list(sps).index('AIA:3(49)') + 1          # inserted after Art 3
    assert 'Article 4a(1)' in sps['AIA:10(1)'].text and sps['AIA:10(1)'].amended_by
    assert 'for the purposes of this definition' in sps['AIA:3(14)'].text
    raw = read_model(os.path.join(ROOT, 'packs', 'AIA', 'sources', 'AIA', 'raw.json'), model.RawDocument)
    orig = {p.ref: p for p in raw.provisions}
    assert '10(5)(f)' in orig and 'for the purposes' not in orig['3(14)'].text


def test_mr_amendments_applied(cat):
    sps = cat.sps
    order = list(sps)
    assert order.index('MR:8-3') == order.index('MR:8-2') + 1 and order.index('MR:8-4') == order.index('MR:8-3') + 1
    assert order.index('MR:20(10)') == order.index('MR:20(9)') + 1
    assert sps['MR:8-3'].amended_by == ['Art 3(1) Reg. (EU) 2026/1744']
    assert sps['MR:54-2'].text == 'It shall apply from 20 January 2027.' and sps['MR:54-2'].amended_by
    assert cat.packs['MR'][1].amending_acts == ['MR-COR', 'OMN-MR'] and cat.packs['AIA'][1].amending_acts == ['OMN']


def test_link_ids_generated_in_catalogue_order(cat):
    links = cat.links()
    assert [l[0] for l in links[:3]] == ['L0001', 'L0002', 'L0003'] and links[-1][0] == 'L0384'
    assert links[0][1] == next(iter(cat.srs))


def test_licensed_text_not_exported():
    cat = load_catalogue(SAMPLE, [DEMO])
    demo = lambda rows: [r for r in rows if r['sd_id'] == 'DEMO']
    assert demo(to_tables(cat)['SP']) and all(not r['text'] for r in demo(to_tables(cat)['SP']))
    assert all(r['text'] for r in demo(to_tables(cat, internal=True)['SP']))


# ------------------------------------------------------------------------------------------------ packs (spec § 8.3)
def test_packs_loaded_in_known_order():
    cat = load_catalogue(SAMPLE, [EXTRA])                # a directory of packs
    E, W, st = checks.check(cat)
    assert E == [] and W == ['SR SR.AIA.0005: coverage incomplete -> reported as Gap']
    assert list(cat.packs) == ['MR', 'AIA', 'MR-AIDA', 'EN18228', 'TR5469', 'DEMO'] and cat.missing_packs() == []
    assert (st['PACKS'], st['UNREACHED'], st['COND']) == (6, 0, 1)
    assert cat.pack_of_sd['DEMO'] == 'DEMO' and [d.dec_id for d in cat.decisions][-1] == 'D9001'
    pub = load_catalogue(SAMPLE, [])
    assert 'DEMO' not in pub.sds and pub.missing_packs() == ['DEMO']
    E, W, st = checks.check(pub)
    assert E == [] and st['UNREACHED'] == 1 and not any('CR.PERF.901' in w for w in W)   # only DEMO reaches it
    assert list(load_catalogue(SAMPLE, [], defaults=False).packs) == []


def test_orphans_warned_when_all_known_packs_loaded(sample):
    edit(sample, 'common_requirements.json', lambda d: d['requirements'].append(
        dict(cr_id='CR.PERF.902', title='Not linked', statement='Demo.', cr_type='TEST', object_level='AIF',
             acceptance_criterion='x', status='DRAFT')))
    E, W, st = checks.check(load_catalogue(sample, []))                 # DEMO missing: counted, not warned
    assert E == [] and st['UNREACHED'] == 2 and not any('CR.PERF.902' in w for w in W)
    E, W, st = checks.check(load_catalogue(sample, [DEMO]))
    assert st['UNREACHED'] == 0 and any(w.startswith('CR CR.PERF.902: no counting SR link') for w in W)


def test_pack_reference_rules(sample, tmp_path):
    demo = os.path.join(tmp_path, 'DEMO')
    shutil.copytree(DEMO, demo)
    edit(sample, '../packs/EN18228/pack.json', lambda d: d.pop('requires'))          # ZA correspondence to AIA
    edit(demo, 'pack.json', lambda d: d.update(requires=['other']))
    E = checks.check(load_catalogue(sample, [demo]))[0]
    for x in ['CORR C0001: SP AIA:9(8) of pack AIA not visible from pack EN18228', 'PACK DEMO: requires pack other']:
        assert any(x in e for e in E), (x, E)
    edit(demo, 'pack.json', lambda d: (d.pop('requires'), d.update(pack_id='DEMO2')))
    W = checks.check(load_catalogue(sample, [demo]))[1]
    assert 'PACK DEMO2: not a known pack of the master list (catalogue.json known_packs)' in W


def test_pack_and_source_ids_unique(sample, tmp_path):
    a, b = os.path.join(tmp_path, 'a'), os.path.join(tmp_path, 'b')
    shutil.copytree(DEMO, a); shutil.copytree(DEMO, b)
    assert any('PACK DEMO: loaded twice' in e for e in checks.check(load_catalogue(sample, [a, b]))[0])
    edit(b, 'pack.json', lambda d: d.update(pack_id='DEMO2'))
    assert 'SD DEMO: defined in packs DEMO and DEMO2' in checks.check(load_catalogue(sample, [a, b]))[0]


def test_amending_act_only_on_own_pack(tmp_path):
    root = os.path.join(tmp_path, 'repo')
    shutil.copytree(CAT, os.path.join(root, 'catalogue')); shutil.copytree(os.path.join(ROOT, 'packs'), os.path.join(root, 'packs'))
    edit(root, 'packs/AIA/sources/OMN/raw.json', lambda d: d.update(amends='MR'))
    E = checks.check(load_catalogue(os.path.join(root, 'catalogue'), []))[0]
    assert 'OMN: amends MR, which is not a raw document of pack AIA' in E


def test_resolve_packs(tmp_path, monkeypatch):
    repo = os.path.join(tmp_path, 'repo')
    root = os.path.join(repo, 'catalogue')
    os.makedirs(root)
    json.dump({'baseline': 'X', 'known_packs': [{'pack_id': 'P1', 'access': 'public', 'path': '../packs/P1'},
                                                {'pack_id': 'Q', 'access': 'private'}]},
              open(os.path.join(root, 'catalogue.json'), 'w'))
    for d in ('private/packs/Q', 'private/packs/R'):
        os.makedirs(os.path.join(tmp_path, d)); open(os.path.join(tmp_path, d, 'pack.json'), 'w').write('{}')
    monkeypatch.delenv('CTM_PACKS', raising=False)
    p1 = os.path.join(repo, 'packs', 'P1')
    assert resolve_packs(root) == [p1] == default_packs(root)
    json.dump({'packs': ['../private/packs']}, open(os.path.join(repo, 'ctm.local.json'), 'w'))
    priv = [os.path.join(str(tmp_path), 'private', 'packs', x) for x in 'QR']
    assert resolve_packs(root) == [p1] + priv                           # a directory of packs is expanded
    monkeypatch.setenv('CTM_PACKS', os.pathsep.join(['a', 'b']))
    assert resolve_packs(root) == [p1, os.path.abspath('a'), os.path.abspath('b')]
    assert resolve_packs(root, []) == [p1] and resolve_packs(root, [], defaults=False) == []
    assert resolve_packs(root, [priv[0]]) == [p1, priv[0]]


def test_load_rejects_a_pack_as_root():
    with pytest.raises(CatalogueError, match='is a pack'):
        load_catalogue(DEMO, [])


# ------------------------------------------------------------------------------------------------ link conditions
def build_cis(tmp_path, root, packs, basis=None, tags=None, name='p'):
    """Build the sample project; {cr_id: (applicability, justification, guidance)} of its CIs (first object)."""
    import build_project, make_catalogue
    from openpyxl import load_workbook
    cat = os.path.join(tmp_path, f'{name}_cat.xlsx')
    make_catalogue.build(cat, to_tables(load_catalogue(root, packs)))
    s = load_project_seed(SEEDS[0]).model_copy(deep=True)
    s.tags.update(tags or {})
    if basis is not None:
        s.basis = [model.BasisEntry(sd_id=x) for x in basis]
    prj = os.path.join(tmp_path, f'{name}.xlsx')
    build_project.init_project(cat, prj, s)
    info = build_project.build_project(cat, prj)
    ws = load_workbook(prj)['CI']
    h = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            out.setdefault(r[h.index('CR ID')], (r[h.index('Applicability')], r[h.index('Applicability justification')] or '',
                                                 r[h.index('Guidance')] or ''))
    return out, info


APPL, NA = vocab.label_of('APPL', 'APPLICABLE'), vocab.label_of('APPL', 'NA_FEAT')
ALL = ['MR', 'AIA', 'MR-AIDA', 'EN18228', 'TR5469']


def test_link_condition_or_across_sources(tmp_path):
    """CR.FS.001 is linked by MR (no condition) and by DEMO (condition AI.usage_level = A1)."""
    b1 = {'AI.usage_level': 'B1'}
    ci, _ = build_cis(tmp_path, SAMPLE, [DEMO], ALL + ['DEMO'], b1, 'all')
    assert ci['CR.FS.001'][0] == APPL and '[DEMO] Demo rating' in ci['CR.FS.001'][2]       # MR requires it anyway
    ci, _ = build_cis(tmp_path, SAMPLE, [DEMO], ['AIA', 'DEMO'], b1, 'demo_b1')
    assert ci['CR.FS.001'][:2] == (NA, 'Rule not met: (sys.control_system = yes) AND (AI.usage_level = A1)')
    ci, _ = build_cis(tmp_path, SAMPLE, [DEMO], ['AIA', 'DEMO'], {'AI.usage_level': 'A1'}, 'demo_a1')
    assert ci['CR.FS.001'][0] == APPL
    ci, _ = build_cis(tmp_path, SAMPLE, [DEMO], ['AIA', 'DEMO'], {'AI.usage_level': ''}, 'demo_unset')
    assert ci['CR.FS.001'][0] == APPL and 'undetermined (AI.usage_level)' in ci['CR.FS.001'][1]


def test_nearest_link_of_a_source_decides(sample, tmp_path):
    """DEMO links parent CR.ROB.001 without condition and its child CR.ROB.001.1 with a condition: for ROB.001.1 the
    child link decides; ROB.001.2 is decided by the parent link. MR's unconditional link on ROB.001 still counts
    (AIA also links ROB.001, so the bases leave it out)."""
    demo = os.path.join(tmp_path, 'DEMO')
    shutil.copytree(DEMO, demo)
    edit(demo, 'sources/DEMO/derived.json', lambda d: d['requirements'][1]['links'].extend([
        dict(cr='CR.ROB.001', type='COVERS'), dict(cr='CR.ROB.001.1', type='COVERS', condition='AI.usage_level = A1')]))
    b1 = {'AI.usage_level': 'B1'}
    ci, _ = build_cis(tmp_path, sample, [demo], ['DEMO'], b1, 'near')
    assert ci['CR.ROB.001.1'][0] == NA and ci['CR.ROB.001.2'][0] == APPL
    ci, _ = build_cis(tmp_path, sample, [demo], ['MR', 'DEMO'], b1, 'near_mr')
    assert ci['CR.ROB.001.1'][0] == APPL


def test_link_conditions_checked(sample, tmp_path):
    demo = os.path.join(tmp_path, 'DEMO')
    shutil.copytree(DEMO, demo)

    def links(d):
        d['requirements'][0]['links'][0].update(condition='AI.usage_level = Z9')            # PARTLY: counting
        d['requirements'][0]['links'][1].update(condition='(AI.learning = static')
        d['requirements'][1]['links'].append(dict(cr='CR.PMM.001', type='CONTEXT', condition='AI.learning = static'))
    edit(demo, 'sources/DEMO/derived.json', links)
    E, W, _ = checks.check(load_catalogue(sample, [demo]))
    assert any('SR.DEMO.0001 -> CR.PERF.004.3 condition: rule value "Z9" not allowed for AI.usage_level' in e for e in E), E
    assert any('SR.DEMO.0001 -> CR.PERF.901 condition: rule syntax' in e for e in E), E
    assert 'SR SR.DEMO.0002 -> CR.PMM.001: condition on a non-counting link has no effect' in W


# ------------------------------------------------------------------------------------------------ refinement sets
def test_refinement_sets(sample, tmp_path):
    """CR.LOG.001 is fully refined by CR.PERF.004.2 and CR.PERF.004.3. With PERF.004.3 not applicable, one set
    {PERF.004.2, PERF.004.3} is incomplete (LOG.001 assessed); two sets {PERF.004.2} and {PERF.004.3} substitute."""
    def rules_(d):
        by = {c['cr_id']: c for c in d['requirements']}
        by['CR.PERF.004.3']['applicability_rule'] = 'AI.learning = dynamic'
    edit(sample, 'common_requirements.json', rules_)
    static = {'AI.learning': 'static'}
    ci, info = build_cis(tmp_path, sample, [], ALL, static, 'one_set')
    assert 'CR.LOG.001' in ci and ci['CR.PERF.004.3'][0] == NA

    def sets(d):
        by = {c['cr_id']: c for c in d['requirements']}
        by['CR.PERF.004.2']['refines'][0]['set'] = 'A'
        by['CR.PERF.004.3']['refines'][0]['set'] = 'B'
    edit(sample, 'common_requirements.json', sets)
    cat = load_catalogue(sample, [])
    assert {r['cr_id']: r['refines_full'] for r in to_tables(cat)['CR']}['CR.PERF.004.2'] == 'CR.LOG.001 [A]'
    assert checks.check(cat)[0] == []
    ci, info2 = build_cis(tmp_path, sample, [], ALL, static, 'two_sets')
    assert 'CR.LOG.001' not in ci and info2['substituted'] == info['substituted'] + 1


def test_set_on_partial_refinement_warned(sample):
    edit(sample, 'common_requirements.json', lambda d: d['requirements'][-1]['refines'][0].update(set='X'))
    W = checks.check(load_catalogue(sample, []))[1]
    assert 'CR CR.PERF.901: refinement set "X" on a partial refinement of CR.PERF.004.3 has no effect' in W


# ------------------------------------------------------------------------------------------------ files and schemas
ROOTS = [CAT, *default_packs(CAT), SAMPLE, *default_packs(SAMPLE), DEMO]


@pytest.mark.parametrize('root', ROOTS)
def test_files_canonical(root):
    """load -> save is byte-identical (fix with: python check_catalogue.py DIR --fmt)."""
    for path, cls in catalogue_files(root):
        assert dumps_model(read_model(path, cls), path) == open(path, encoding='utf-8').read(), path


@pytest.mark.parametrize('path', SEEDS)
def test_seeds_canonical(path):
    assert dumps_model(load_project_seed(path), path) == open(path, encoding='utf-8').read()


def test_schemas_up_to_date():
    for fn, txt in schemas.render().items():
        assert open(os.path.join(ROOT, 'schemas', fn), encoding='utf-8').read() == txt, f'{fn}: run python -m ctm.schemas'
    assert sorted(os.listdir(os.path.join(ROOT, 'schemas'))) == sorted(schemas.render())


def test_files_valid_against_json_schema():
    jsonschema = pytest.importorskip('jsonschema')
    files = [f for r in ROOTS for f in catalogue_files(r)] + [(p, model.ProjectSeed) for p in SEEDS]
    for path, cls in files:
        data = json.load(open(path, encoding='utf-8'))
        sch = json.load(open(os.path.normpath(os.path.join(os.path.dirname(path), data['$schema'])), encoding='utf-8'))
        jsonschema.validate(data, sch)


# ------------------------------------------------------------------------------------------------ Excel layout
SHEET_MODEL = {'SD': model.SourceDocument, 'SP': model.ProvisionAnnotation, 'SR': model.SourceRequirement,
               'CR': model.CommonRequirement, 'LINK_SR_CR': model.CoverageLink, 'CORR': model.Correspondence,
               'DECISIONS': model.Decision, 'BASIS': model.BasisEntry, 'OBJECTS': model.AssessmentObject,
               'EV': model.EvidenceItem}
RENAMED = {('LINK_SR_CR', 'cr_id'): 'cr', ('LINK_SR_CR', 'link_type'): 'type', ('SP', 'native_ref'): 'ref'}


def vocab_of(field):
    def walk(t):
        for m in getattr(t, '__metadata__', ()):
            if isinstance(m, VocabRef): return m.lst
        for a in typing.get_args(t):
            r = walk(a)
            if r: return r
    for m in field.metadata:
        if isinstance(m, VocabRef): return m.lst
    return walk(field.annotation)


@pytest.mark.parametrize('sheet', SHEET_MODEL)
def test_excel_columns_match_model(sheet):
    cls = SHEET_MODEL[sheet]
    cols = {**CATALOGUE, **PROJECT_INPUT}[sheet]
    for c in cols:
        f = cls.model_fields.get(RENAMED.get((sheet, c.key), c.key))
        if f is None:
            continue                                  # export-only column (IDs, sd_id, seq, parent_sp, domain …)
        if c.lst == 'YN':
            assert f.annotation is bool, (sheet, c.key)
        else:
            assert vocab_of(f) == c.lst, (sheet, c.key, vocab_of(f), c.lst)
    bound = {RENAMED.get((sheet, c.key), c.key) for c in cols}
    for name, f in cls.model_fields.items():
        if vocab_of(f):
            assert name in bound, (sheet, name)


def test_tables_have_exactly_the_workbook_columns(cat):
    for sheet, rows in to_tables(cat).items():
        keys = {c.key for c in CATALOGUE[sheet]}
        for r in rows[:1]:
            assert set(r) == keys, sheet


# ------------------------------------------------------------------------------------------------ seeded errors
@pytest.fixture
def sample(tmp_path):
    """Copy of the sample (master list + public packs); returns the master list directory."""
    root = os.path.join(tmp_path, 'sample')
    shutil.copytree(os.path.dirname(SAMPLE), root)
    return os.path.join(root, 'catalogue')


def errors(root):
    return checks.check(load_catalogue(root, []))[0]


def test_sample_clean(sample):
    E, W, _ = checks.check(load_catalogue(sample, []))
    assert E == [] and W == ['SR SR.AIA.0005: coverage incomplete -> reported as Gap']


def test_check_catches_reference_and_rule_errors(sample):
    def mr(d):
        d['requirements'][0]['links'].append({'cr': 'CR.GEN.999', 'type': 'COVERS'})
        d['requirements'][1]['sp_refs'] = ['MR:III-9.9.9']
        d['requirements'][2]['sp_refs'] = ['AIA:12(1)']
        d['requirements'][3]['links'].append(dict(d['requirements'][3]['links'][0]))
    edit(sample, '../packs/MR/sources/MR/derived.json', mr)

    def cr(d):
        d['requirements'][5]['applicability_rule'] = 'sys.control_system = maybe'
        d['requirements'][0].pop('decomposition')
    edit(sample, 'common_requirements.json', cr)
    E = errors(sample)
    for x in ['link to unknown CR CR.GEN.999', 'unknown SP MR:III-9.9.9', 'AIA:12(1) belongs to another SD',
              'duplicate link', 'not allowed for sys.control_system', 'no decomposition']:
        assert any(x in e for e in E), (x, E)


def test_refines_checked(sample):
    def cr(d):
        ids = [c['cr_id'] for c in d['requirements']]
        by = {c['cr_id']: c for c in d['requirements']}
        ref = lambda *ids: [dict(cr=x) for x in ids]
        by['CR.FS.001']['refines'] = ref('CR.ROB.001')                    # valid, non-hierarchical
        by['CR.LOG.001']['refines'] = ref('CR.GEN.999')
        by['CR.PMM.001']['refines'] = ref('CR.PMM.001')
        by['CR.PERF.004.1']['refines'] = ref('CR.PERF.004')
        by['CR.HO.001']['refines'] = ref('CR.PHYS.001')
        by['CR.PHYS.001']['refines'] = ref('CR.HO.001')
        by['CR.LOG.001']['refines'].append(dict(cr='CR.ROB.001', coverage='FULL'))   # full refinement of a parent
        assert 'CR.ROB.001' in ids
    edit(sample, 'common_requirements.json', cr)
    cat = load_catalogue(sample, [])
    E, W, _ = checks.check(cat)
    assert not any('CR.FS.001' in e for e in E + W)
    for x in ['CR.LOG.001: refines unknown CR CR.GEN.999', 'CR.PMM.001: refines itself', 'CR.HO.001: refines cycle',
              'CR.PHYS.001: refines cycle']:
        assert x in ' | '.join(E), (x, E)
    assert any('CR.PERF.004.1: refines its ancestor CR.PERF.004' in w for w in W)
    assert any('CR.LOG.001: full refinement of parent CR CR.ROB.001' in w for w in W)
    assert {r['cr_id']: r['refines'] for r in to_tables(cat)['CR']}['CR.FS.001'] == 'CR.ROB.001'


def test_full_refinement_chain_warned(sample):
    def cr(d):
        by = {c['cr_id']: c for c in d['requirements']}
        by['CR.FS.001']['refines'] = [dict(cr='CR.PERF.004.3', coverage='FULL')]    # PERF.004.3 fully refines LOG.001
    edit(sample, 'common_requirements.json', cr)
    E, W, _ = checks.check(load_catalogue(sample, []))
    assert E == []
    assert any(w.startswith('CR CR.PERF.004.3: full refinement of CR.LOG.001 is chained') for w in W), W


def test_join_errors(sample):
    edit(sample, '../packs/MR/sources/MR/derived.json', lambda d: d['provisions'].pop(0))
    edit(sample, '../packs/AIA/sources/AIA/derived.json', lambda d: d['provisions'].append(
        {'ref': '99(9)', 'role': 'ATOMIC', 'in_catalogue_scope': True}))
    E = errors(sample)
    assert any('MR:III-A(e): provision has no annotation' in e for e in E)
    assert any('AIA:99(9): annotation without provision' in e for e in E)


@pytest.mark.parametrize('rel,mutate,msg', [
    ('../packs/MR/sources/MR/derived.json', lambda d: d['requirements'][0].update(coverage_complete='perhaps'), 'coverage_complete'),
    ('../packs/MR/sources/MR/derived.json', lambda d: d['provisions'][0].update(role='Atomic'), 'register list SPROLE'),
    ('common_requirements.json', lambda d: d['requirements'][0].update(cr_id='CR.XYZ.001'), 'register list DOMAIN'),
    ('common_requirements.json', lambda d: d['requirements'][0].update(colour='red'), 'Extra inputs'),
    ('common_requirements.json', lambda d: d['requirements'][0].update(provenance=['MR']), 'Extra inputs'),
    ('../packs/MR/sources/MR/derived.json', lambda d: d['requirements'][0].update(sr_id='SR-MR-1'), 'sr_id'),
    ('../packs/MR/pack.json', lambda d: d.update(sources=[]), 'sources'),
])
def test_schema_errors_rejected_on_load(sample, rel, mutate, msg):
    edit(sample, rel, mutate)
    with pytest.raises(CatalogueError, match=msg):
        load_catalogue(sample, [])


def test_bad_amendment_target(tmp_path):
    root = os.path.join(tmp_path, 'repo')
    shutil.copytree(CAT, os.path.join(root, 'catalogue')); shutil.copytree(os.path.join(ROOT, 'packs'), os.path.join(root, 'packs'))
    edit(root, 'packs/AIA/sources/OMN/raw.json', lambda d: d['amendments'][0].update(target_ref='3(99)'))
    assert any('target 3(99) not in AIA' in e for e in errors(os.path.join(root, 'catalogue')))


def test_seed_tag_validation():
    seed = json.load(open(SEEDS[0], encoding='utf-8'))
    seed['tags']['AI.learning'] = 'sometimes'
    with pytest.raises(Exception, match='not allowed for tag AI.learning'):
        model.ProjectSeed.model_validate(seed)


def test_vocab_keys_unique():
    for lst, entries in vocab.DEFAULT.items():
        assert len({e[0] for e in entries}) == len(entries), lst
