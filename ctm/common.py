"""Shared styling and sheet I/O helpers."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

from .schema import HEADER_FILL, ACTOR_FILL

F = 'Arial'
HDR_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
HDR_DARK_FONT = Font(name=F, bold=True, color='1F1F1F', size=10)
CHK_FONT = Font(name=F, italic=True, color='595959', size=9)
BODY = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, labels, widths=None, fills=None, row=1, dark=None):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab)
        fill = (fills or {}).get(j)
        if fill:
            c.fill = PatternFill('solid', fgColor=fill)
            c.font = HDR_DARK_FONT
        else:
            c.fill = PatternFill('solid', fgColor=HEADER_FILL)
            c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[row].height = 32


def body_cell(ws, r, c, v, font=None, wrap=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font or BODY
    if wrap:
        cell.alignment = WRAP
    cell.border = BORDER
    return cell


def add_list_validation(ws, lst, rng):
    dv = DataValidation(type='list', formula1=f'L_{lst}', allow_blank=True, showErrorMessage=True,
                        errorTitle='Not in register', error=f'Choose a value from REGISTER list {lst}.')
    ws.add_data_validation(dv)
    dv.add(rng)


def add_status_colours(ws, rng, first_cell, reg, lists):
    """Conditional fills from REGISTER colours; formulas compare against named cells (no literals)."""
    for lst in lists:
        for e in reg[lst]:
            if e['colour'] and e['colour'].upper() != 'FFFFFF':
                ws.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'{first_cell}={lst}_{e["key"]}'], fill=PatternFill('solid', fgColor=e['colour'])))


def read_table(wb, sheet, cols):
    """Read an entity sheet by header labels -> list of dicts keyed by Col.key (rows with empty first column skipped)."""
    ws = wb[sheet]
    hdr = [c.value for c in ws[1]]
    idx = {}
    for col in cols:
        if col.label in hdr:
            idx[col.key] = hdr.index(col.label)
    out = []
    first = cols[0].key
    for row in ws.iter_rows(min_row=2, values_only=True):
        if first not in idx or row[idx[first]] in (None, ''):
            continue
        d = {}
        for col in cols:
            v = row[idx[col.key]] if col.key in idx and idx[col.key] < len(row) else None
            d[col.key] = '' if v is None else (v.strip() if isinstance(v, str) else v)
        out.append(d)
    return out


def col_letter(cols, key):
    for j, c in enumerate(cols, 1):
        if c.key == key:
            return get_column_letter(j)
    raise KeyError(key)


def combine_formula(F_, G, X, O, P, N, R, empty):
    """Roll-up precedence: Fail > Gap > Not assessed > Open > Pass > N/A; `empty` if nothing counted.
    Arguments are formula fragments evaluating to counts. R = dict of ROLLUP named cells."""
    return (f'IF(({F_})>0,{R["FAIL"]},IF(({G})>0,{R["GAP"]},IF(({X})>0,{R["NOTASS"]},'
            f'IF(({O})>0,{R["OPEN"]},IF(({P})>0,{R["PASS"]},IF(({N})>0,{R["NA"]},{empty}))))))')


R = {k: f'ROLLUP_{k}' for k in ['PASS', 'FAIL', 'GAP', 'NOTASS', 'OOS', 'OPEN', 'NA', 'NOTINST', 'NOREQ']}


def counts_over(status_rng, crit_prefix):
    """Return dict of COUNTIFS fragments for each roll-up bucket over a status range, with extra criteria prefix."""
    def c(name):
        return f'COUNTIFS({crit_prefix}{status_rng},{R[name]})'
    # NOTINST and NOREQ are not counted: if nothing is counted the caller's `empty` fallback applies
    return dict(F=c('FAIL'), G=c('GAP'), X=f"{c('NOTASS')}+{c('OOS')}", O=c('OPEN'), P=c('PASS'), N=c('NA'))
