#!/usr/bin/env python3
"""Create or rebuild a project workbook from a catalogue (spec § 5.4–5.6).

  python build_project.py CATALOGUE.xlsx PROJECT.xlsx --init [--seed SEED.json]   # create input sheets
  python build_project.py CATALOGUE.xlsx PROJECT.xlsx                                # (re)build CIs and views

Rebuild keeps all editable CI fields (matched by CI ID = <object>/<CR>), keeps manual applicability
overrides (applicability differs from the previous rule result), and moves CIs that no longer exist
to sheet CI_ORPHANED. A backup PROJECT.xlsx.bak is written before overwriting.
"""
import argparse, os, re, shutil, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ctm.schema import (CATALOGUE, PROJECT_INPUT, PRJ_FIELDS, CI_COLS, CI_EDITABLE, ACTOR_FILL,
                        CR_MAX_DEPTH, SP_MAX_DEPTH)
from ctm.register import write_register, read_register, as_default_format, labels, label
from ctm.common import (header, body_cell, add_list_validation, add_status_colours, read_table, col_letter,
                        combine_formula, counts_over, R, BODY, BOLD, CHK_FONT, F, WRAP)
from ctm.project import AUTO_PREFIX, natkey, select_project, checklist_rows, split_full, split_refs  # noqa: F401
from ctm.store import load_project_seed, seed_inputs


# ------------------------------------------------------------------------------------------ init
def init_project(cat_path, prj_path, seed=None):
    """Create the project input sheets; `seed` (ctm.model.ProjectSeed) pre-fills them."""
    cat = load_workbook(cat_path)
    reg = read_register(cat)
    wb = Workbook()
    wb.remove(wb.active)
    write_register(wb, as_default_format(reg))
    if seed:
        prj, basis, tags, objects, ev = seed_inputs(seed)
    else:
        prj, tags, objects, ev = {}, {}, [], []
        basis = [dict(sd_id=r['sd_id'], edition=r['edition']) for r in read_table(cat, 'SD', CATALOGUE['SD'])]
    _write_inputs(wb, reg, prj=prj, basis=basis, tags=tags, objects=objects, ev=ev)
    wb.move_sheet('REGISTER', offset=len(wb.sheetnames))
    wb.save(prj_path)
    return prj_path


def _write_inputs(wb, reg, prj, basis, tags, objects, ev):
    ws = wb.create_sheet('PRJ', 0)
    header(ws, ['Field', 'Value'], [34, 90], {1: ACTOR_FILL['assessor'], 2: ACTOR_FILL['assessor']})
    for i, f in enumerate(PRJ_FIELDS, 2):
        body_cell(ws, i, 1, f, BOLD); body_cell(ws, i, 2, prj.get(f, ''))
    sheets = {'BASIS': basis, 'OBJECTS': objects, 'EV': ev}
    for name in ['BASIS', 'TAGS', 'OBJECTS', 'EV']:
        cols = PROJECT_INPUT[name]
        ws = wb.create_sheet(name)
        header(ws, [c.label for c in cols], [c.width for c in cols], {j: ACTOR_FILL[c.actor] for j, c in enumerate(cols, 1)})
        if name == 'TAGS':
            for i, e in enumerate(reg['TAG'], 2):
                grp, _, desc = (e['desc'] or '').partition(' | ')
                body_cell(ws, i, 1, e['label']); body_cell(ws, i, 2, grp)
                body_cell(ws, i, 3, tags.get(e['label'], '') or None); body_cell(ws, i, 4, desc)
                dv = DataValidation(type='list', formula1=f'L_{e["attr"]}', allow_blank=True)
                ws.add_data_validation(dv); dv.add(f'C{i}')
        else:
            for i, rec in enumerate(sheets[name], 2):
                for j, c in enumerate(cols, 1):
                    body_cell(ws, i, j, rec.get(c.key, '') or None)
            for j, c in enumerate(cols, 1):
                if c.lst:
                    L = get_column_letter(j); add_list_validation(ws, c.lst, f'{L}2:{L}500')
        ws.freeze_panes = 'B2'


# ------------------------------------------------------------------------------------------ build
def build_project(cat_path, prj_path, out_path=None):
    out_path = out_path or prj_path
    cat = load_workbook(cat_path)
    reg = read_register(cat)
    C = {s: read_table(cat, s, cols) for s, cols in CATALOGUE.items()}
    old = load_workbook(prj_path)
    prj = {r[0]: (r[1] or '') for r in old['PRJ'].iter_rows(min_row=2, values_only=True) if r and r[0]}
    P = {s: read_table(old, s, cols) for s, cols in PROJECT_INPUT.items()}
    tags = {r['tag']: str(r['value']) if r['value'] not in ('', None) else '' for r in P['TAGS']}
    old_ci = {}
    if 'CI' in old.sheetnames:
        for r in read_table(old, 'CI', CI_COLS):
            old_ci[r['ci_id']] = r
    pid = prj.get('Project ID') or 'PRJ'

    role_attr = {e['label']: e['attr'] for e in reg['SPROLE']}
    sd_status = {r['sd_id']: r['lifecycle_status'] for r in C['SD']}
    basis = [b['sd_id'] for b in P['BASIS'] if b['sd_id']]
    objects = P['OBJECTS']
    S = select_project(C, reg, basis, tags, objects)
    sp, sr, links, crs, children, corr, subst, guid = S.sp, S.sr, S.links, S.crs, S.children, S.corr, S.subst, S.guid

    # ---- merge with the previous CIs: keep project entries and manual applicability overrides
    ci_rows = []
    for row in S.ci_rows:
        prev = old_ci.pop(row['ci_id'], None)
        if prev:
            manual = prev['applicability'] not in ('', None) and prev['applicability'] != prev['rule_applicability']
            for f in CI_EDITABLE:
                if f in ('applicability', 'applicability_justification'):
                    continue
                row[f] = prev.get(f, '')
            if manual:
                row['applicability'] = prev['applicability']
                row['applicability_justification'] = prev['applicability_justification']
            elif prev['applicability_justification'] and not str(prev['applicability_justification']).startswith(AUTO_PREFIX):
                row['applicability_justification'] = prev['applicability_justification']
        ci_rows.append(row)
    orphaned = list(old_ci.values())

    # ---- workbook
    wb = Workbook(); wb.remove(wb.active)
    write_register(wb, as_default_format(reg))
    _write_inputs(wb, reg, prj, P['BASIS'], tags, objects, P['EV'])
    nCI = max(len(ci_rows), 1) + 1
    ci = _write_ci(wb, reg, ci_rows, nCI)
    if orphaned:
        ws = wb.create_sheet('CI_ORPHANED')
        header(ws, [c.label for c in CI_COLS], [c.width for c in CI_COLS])
        for i, r in enumerate(orphaned, 2):
            for j, c in enumerate(CI_COLS, 1):
                body_cell(ws, i, j, r.get(c.key) or None)
    _ev_counts(wb, nCI)
    v_cr = _v_cr(wb, crs, children, nCI, subst)
    v_link = _v_link(wb, links, crs, v_cr)
    v_sr = _v_sr(wb, sr, v_link)
    v_sp = _v_sp(wb, sp, role_attr, v_sr)
    _v_presumption(wb, corr, sd_status, v_sp)
    _v_register_view(wb, 'V_EXCLUSIONS', 'excl_idx', nCI,
                     ['ci_id', 'object_id', 'cr_id', 'title', 'applicability_justification'],
                     'Items not in scope (agreed) – to be disclosed in the conformity report')
    _v_register_view(wb, 'V_FINDINGS', 'find_idx', nCI,
                     ['ci_id', 'object_id', 'cr_id', 'title', 'verdict', 'finding_severity', 'assessor_remarks'],
                     'Findings register (Major / Minor / Observation)')
    for sd_id in basis:
        _v_checklist(wb, S, sd_id, ci_rows, v_sr)
    _v_summary(wb, reg, prj, basis, nCI, v_cr, v_sr, v_sp)
    _snapshot(wb, C)
    _readme(wb)
    order = ['README', 'PRJ', 'BASIS', 'TAGS', 'OBJECTS', 'CI', 'EV', 'V_SUMMARY', 'V_FINDINGS', 'V_EXCLUSIONS']
    order += [s for s in wb.sheetnames if s.startswith('V_CL_')]
    order += ['V_SP', 'V_SR', 'V_LINK', 'V_CR', 'V_PRESUMPTION']
    order += [s for s in wb.sheetnames if s not in order]
    wb._sheets = [wb[s] for s in order]
    if os.path.exists(out_path) and out_path == prj_path:
        shutil.copyfile(out_path, out_path + '.bak')
    wb.save(out_path)
    return dict(ci=len(ci_rows), cr=len(crs), sr=len(sr), sp=len(sp), links=len(links), orphaned=len(orphaned),
                substituted=len(subst))


# ------------------------------------------------------------------------------------------ CI
def _write_ci(wb, reg, rows, n):
    ws = wb.create_sheet('CI')
    cols = CI_COLS
    header(ws, [c.label for c in cols], [c.width for c in cols], {j: ACTOR_FILL[c.actor] for j, c in enumerate(cols, 1)})
    L = {c.key: get_column_letter(j) for j, c in enumerate(cols, 1)}
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            if c.key in ('status', 'excl_idx', 'find_idx'):
                continue
            body_cell(ws, i, j, r.get(c.key) if r.get(c.key) not in ('', None) else None)
        a, v, s = L['applicability'], L['verdict'], L['finding_severity']
        body_cell(ws, i, cols.index(next(c for c in cols if c.key == 'status')) + 1,
                  f'=IF({a}{i}="","",IF({a}{i}=APPL_APPLICABLE,IF({v}{i}="",VERDICT_PEND,{v}{i}),{a}{i}))')
        ws[f'{L["excl_idx"]}{i}'] = f'=IF({a}{i}=APPL_NIS,COUNTIF(${a}$2:{a}{i},APPL_NIS),"")'
        ws[f'{L["find_idx"]}{i}'] = f'=IF({s}{i}="","",COUNTIF(${s}$2:{s}{i},"?*"))'
        ws[f'{L["excl_idx"]}{i}'].font = CHK_FONT; ws[f'{L["find_idx"]}{i}'].font = CHK_FONT
    last = max(n, 2)
    for key, lst in [('applicability', 'APPL'), ('verdict', 'VERDICT'), ('finding_severity', 'SEV')]:
        add_list_validation(ws, lst, f'{L[key]}2:{L[key]}{last}')
    add_status_colours(ws, f'{L["status"]}2:{L["status"]}{last}', f'${L["status"]}2', reg, ['VERDICT', 'APPL'])
    add_status_colours(ws, f'{L["verdict"]}2:{L["verdict"]}{last}', f'${L["verdict"]}2', reg, ['VERDICT'])
    add_status_colours(ws, f'{L["finding_severity"]}2:{L["finding_severity"]}{last}', f'${L["finding_severity"]}2', reg, ['SEV'])
    add_status_colours(ws, f'{L["applicability"]}2:{L["applicability"]}{last}', f'${L["applicability"]}2', reg, ['APPL'])
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'
    for k in ('excl_idx', 'find_idx', 'rule_applicability'):
        ws.column_dimensions[L[k]].hidden = k != 'rule_applicability'
    return dict(n=last, L=L)


def _ci_rng(key, n):
    return f"CI!${col_letter(CI_COLS, key)}$2:${col_letter(CI_COLS, key)}${n}"


def _ev_counts(wb, n):
    ws = wb['EV']
    j = len(PROJECT_INPUT['EV']) + 1
    header_cell = ws.cell(row=1, column=j, value='# CIs citing (chk)')
    header_cell.font = Font(name=F, bold=True, size=10)
    ws.column_dimensions[get_column_letter(j)].width = 12
    for i in range(2, 502):
        c = ws.cell(row=i, column=j, value=f'=IF(A{i}="","",COUNTIF({_ci_rng("evidence_refs", n)},"*"&A{i}&"*"))')
        c.font = CHK_FONT


# ------------------------------------------------------------------------------------------ roll-up views
def _own_buckets(cell):
    """buckets from one status cell (own status) – NOTINST/NOREQ not counted."""
    return dict(F=f'({cell}={R["FAIL"]})*1', G=f'({cell}={R["GAP"]})*1',
                X=f'({cell}={R["NOTASS"]})+({cell}={R["OOS"]})', O=f'({cell}={R["OPEN"]})*1',
                P=f'({cell}={R["PASS"]})*1', N=f'({cell}={R["NA"]})*1')


def _v_cr(wb, crs, children, nCI, subst=None):
    subst = subst or {}
    ws = wb.create_sheet('V_CR')
    S = [f'S{k}' for k in range(CR_MAX_DEPTH + 1)]
    hdr = ['CR ID', 'Domain', 'Title', 'Parent CR', 'Decomposition', 'Object level', '# CIs', 'n Fail', 'n Not in scope',
           'n Open', 'n Pass (incl. Inherited)', 'n N/A'] + S + ['Final status', 'Refines (non-hierarchical)',
                                                                  'Substituted by (full refinements)']
    header(ws, hdr, [16, 8, 34, 16, 9, 8, 7, 7, 8, 7, 9, 7] + [12] * len(S) + [16, 20, 24])
    n = len(crs) + 1
    cicr, cist = _ci_rng('cr_id', nCI), _ci_rng('status', nCI)
    first_s = 13
    row_of = {k: i for i, k in enumerate(crs, 2)}
    for i, (k, r) in enumerate(crs.items(), 2):
        for j, v in enumerate([k, r['domain'], r['title'], r['parent_cr'], r['decomposition'], r['object_level']], 1):
            body_cell(ws, i, j, v or None)
        body_cell(ws, i, 7, f'=COUNTIF({cicr},$A{i})')
        q = lambda name: f'COUNTIFS({cicr},$A{i},{cist},{name})'
        body_cell(ws, i, 8, f'={q("VERDICT_FAIL")}')
        body_cell(ws, i, 9, f'={q("APPL_NIS")}')
        body_cell(ws, i, 10, f'={q("VERDICT_INFO")}+{q("VERDICT_PEND")}')
        body_cell(ws, i, 11, f'={q("VERDICT_PASS")}+{q("APPL_INHERITED")}')
        body_cell(ws, i, 12, f'={q("APPL_NA_FEAT")}+{q("APPL_NA_HAZ")}')
        if k in subst:
            # substituted (D0033): every level S0..Sn combines the same level of its full refinements, referenced cell by
            # cell (no range over the own column, so no circular reference); parents of k read it as a normal child
            for d in range(0, CR_MAX_DEPTH + 1):
                col = get_column_letter(first_s + d)
                cells = [f'{col}{row_of[y]}' for y in subst[k]]
                b = {name: '+'.join(f'({c}={R[key]})*1' for c in cells) for name, key in
                     [('F', 'FAIL'), ('G', 'GAP'), ('O', 'OPEN'), ('P', 'PASS'), ('N', 'NA')]}
                b['X'] = '+'.join(f'({c}={R["NOTASS"]})+({c}={R["OOS"]})' for c in cells)
                body_cell(ws, i, first_s + d, '=' + combine_formula(b['F'], b['G'], b['X'], b['O'], b['P'], b['N'], R,
                                                                    R['NOTINST']))
            body_cell(ws, i, first_s + CR_MAX_DEPTH + 1, f'={get_column_letter(first_s + CR_MAX_DEPTH)}{i}', BOLD)
            body_cell(ws, i, first_s + CR_MAX_DEPTH + 2, r.get('refines') or None)
            body_cell(ws, i, first_s + CR_MAX_DEPTH + 3, '; '.join(subst[k]))
            continue
        body_cell(ws, i, first_s, '=' + combine_formula(f'H{i}', '0', f'I{i}', f'J{i}', f'K{i}', f'L{i}', R, R['NOTINST']))
        s0 = f'{get_column_letter(first_s)}{i}'
        own = _own_buckets(s0)
        for d in range(1, CR_MAX_DEPTH + 1):
            prev = get_column_letter(first_s + d - 1)
            ch = counts_over(f'${prev}$2:${prev}${n}', f'$D$2:$D${n},$A{i},')
            tot = {b: f'{own[b]}+{ch[b]}' for b in own}
            and_f = combine_formula(tot['F'], tot['G'], tot['X'], tot['O'], tot['P'], tot['N'], R, s0)
            f = f'=IF(AND($E{i}=DECOMP_OR,({ch["P"]})>0,{s0}<>{R["FAIL"]}),{R["PASS"]},{and_f})'
            body_cell(ws, i, first_s + d, f)
        body_cell(ws, i, first_s + CR_MAX_DEPTH + 1, f'={get_column_letter(first_s + CR_MAX_DEPTH)}{i}', BOLD)
        body_cell(ws, i, first_s + CR_MAX_DEPTH + 2, r.get('refines') or None)

    fin = get_column_letter(first_s + CR_MAX_DEPTH + 1)
    _rollup_colours(wb, ws, f'{get_column_letter(first_s)}2:{fin}{n}', f'{get_column_letter(first_s)}2')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:{get_column_letter(first_s + CR_MAX_DEPTH + 3)}{n}'
    return dict(n=n, id='A', final=fin)


def _rollup_colours(wb, ws, rng, first_cell):
    """Conditional fills from REGISTER ROLLUP colours; first_cell is relative (top-left of rng)."""
    add_status_colours(ws, rng, first_cell, read_register(wb), ['ROLLUP'])


def _v_link(wb, links, crs, v_cr):
    ws = wb.create_sheet('V_LINK')
    header(ws, ['Link ID', 'SR ID', 'CR ID', 'Link type', 'Counts?', 'CR final status', 'Condition'], [10, 18, 16, 17, 8, 16, 30])
    n = max(len(links), 1) + 1
    vc = f"V_CR!$A$2:$A${v_cr['n']}"; vf = f"V_CR!${v_cr['final']}$2:${v_cr['final']}${v_cr['n']}"
    for i, l in enumerate(links, 2):
        for j, v in enumerate([l['link_id'], l['sr_id'], l['cr_id'], l['link_type']], 1):
            body_cell(ws, i, j, v)
        body_cell(ws, i, 5, f'=IF(COUNTIFS(REG_LIST,"LINKTYPE",REG_LABEL,D{i},REG_ATTR,YN_YES)>0,YN_YES,YN_NO)')
        body_cell(ws, i, 6, f'=IFERROR(INDEX({vf},MATCH(C{i},{vc},0)),{R["NOTINST"]})')
        body_cell(ws, i, 7, l.get('condition') or None)
    _rollup_colours(wb, ws, f'F2:F{n}', 'F2')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:G{n}'
    return dict(n=n)


def _v_sr(wb, sr, v_link):
    ws = wb.create_sheet('V_SR')
    hdr = ['SR ID', 'SD', 'SP refs', '_pipes', 'Statement', 'Obligation', 'Coverage complete', '# counting links',
           'c Fail', 'c Gap', 'c Not assessed', 'c Open', 'c Pass', 'c N/A', 'SR status']
    header(ws, hdr, [18, 9, 26, 6, 50, 10, 9, 8, 6, 6, 8, 6, 6, 6, 16])
    rows = sorted(sr.values(), key=lambda r: (r['sd_id'], natkey(r['sr_id'])))
    n = max(len(rows), 1) + 1
    nl = v_link['n']
    lsr, lcnt, lst = f'V_LINK!$B$2:$B${nl}', f'V_LINK!$E$2:$E${nl}', f'V_LINK!$F$2:$F${nl}'
    for i, r in enumerate(rows, 2):
        vals = [r['sr_id'], r['sd_id'], r['sp_refs'], '|' + '|'.join(r['refs']) + '|', r['statement'], r['obligation'],
                r['coverage_complete']]
        for j, v in enumerate(vals, 1):
            body_cell(ws, i, j, v)
        body_cell(ws, i, 8, f'=COUNTIFS({lsr},$A{i},{lcnt},YN_YES)')
        c = counts_over(lst, f'{lsr},$A{i},{lcnt},YN_YES,')
        for j, b in enumerate(['F', 'G', 'X', 'O', 'P', 'N'], 9):
            body_cell(ws, i, j, f'={c[b]}')
        comb = combine_formula(f'I{i}', f'J{i}', f'K{i}', f'L{i}', f'M{i}', f'N{i}', R, R['NOTINST'])
        body_cell(ws, i, 15, f'=IF(G{i}<>YN_YES,{R["GAP"]},IF(H{i}=0,{R["GAP"]},{comb}))', BOLD)
    ws.column_dimensions['D'].hidden = True
    _rollup_colours(wb, ws, f'O2:O{n}', 'O2')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:O{n}'
    return dict(n=n)


def _v_sp(wb, sp, role_attr, v_sr):
    ws = wb.create_sheet('V_SP')
    S = [f'S{k}' for k in range(SP_MAX_DEPTH + 1)]
    hdr = ['SP ID', 'SD', 'Native ref', 'Heading / text', 'Parent SP', 'Role', '_noreq', '_children', 'In catalogue scope',
           '# own SRs', 'o Fail', 'o Gap', 'o Not assessed', 'o Open', 'o Pass', 'o N/A'] + S + ['Final status', 'Seq']
    header(ws, hdr, [26, 9, 16, 40, 24, 12, 5, 5, 9, 7, 6, 6, 8, 6, 6, 6] + [12] * len(S) + [16, 7])
    rows = sorted(sp.values(), key=lambda r: (r['sd_id'], float(r['seq'] or 0), natkey(r['sp_id'])))
    n = max(len(rows), 1) + 1
    ns = v_sr['n']
    pipes, sst = f'V_SR!$D$2:$D${ns}', f'V_SR!$O$2:$O${ns}'
    s0c = 17
    for i, r in enumerate(rows, 2):
        attr = role_attr.get(r['role'], 'own')
        txt = r['heading'] or (str(r['text'])[:120] if r['text'] else '')
        vals = [r['sp_id'], r['sd_id'], r['native_ref'], txt, r['parent_sp'], r['role'], 1 if attr == 'none' else 0,
                1 if attr == 'children' else 0, r['in_catalogue_scope']]
        for j, v in enumerate(vals, 1):
            body_cell(ws, i, j, v if v != '' else None)
        crit = f'{pipes},"*|"&$A{i}&"|*",'
        body_cell(ws, i, 10, f'=COUNTIF({pipes},"*|"&$A{i}&"|*")')
        c = counts_over(sst, crit)
        for j, b in enumerate(['F', 'G', 'X', 'O', 'P', 'N'], 11):
            body_cell(ws, i, j, f'={c[b]}')
        own = combine_formula(f'K{i}', f'L{i}', f'M{i}', f'N{i}', f'O{i}', f'P{i}', R, R['NOTINST'])
        body_cell(ws, i, s0c, f'=IF(G{i}=1,{R["NOREQ"]},IF(I{i}=YN_NO,{R["OOS"]},IF(J{i}=0,IF(H{i}=1,{R["NOREQ"]},{R["GAP"]}),{own})))')
        s0 = f'{get_column_letter(s0c)}{i}'
        for d in range(1, SP_MAX_DEPTH + 1):
            prev = get_column_letter(s0c + d - 1)
            ch = counts_over(f'${prev}$2:${prev}${n}', f'$E$2:$E${n},$A{i},')
            ob = dict(F=f'K{i}', G=f'L{i}', X=f'M{i}', O=f'N{i}', P=f'O{i}', N=f'P{i}')
            tot = {b: f'{ob[b]}+{ch[b]}' for b in ob}
            comb = combine_formula(tot['F'], tot['G'], tot['X'], tot['O'], tot['P'], tot['N'], R, s0)
            body_cell(ws, i, s0c + d, f'=IF(G{i}=1,{R["NOREQ"]},IF(I{i}=YN_NO,{R["OOS"]},{comb}))')
        fin_c = s0c + SP_MAX_DEPTH + 1
        body_cell(ws, i, fin_c, f'={get_column_letter(fin_c - 1)}{i}', BOLD)
        body_cell(ws, i, fin_c + 1, r['seq'])
    for col in 'GH':
        ws.column_dimensions[col].hidden = True
    fin = get_column_letter(s0c + SP_MAX_DEPTH + 1)
    _rollup_colours(wb, ws, f'{get_column_letter(s0c)}2:{fin}{n}', f'{get_column_letter(s0c)}2')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:{get_column_letter(s0c + SP_MAX_DEPTH + 2)}{n}'
    return dict(n=n, final=fin)


def _v_presumption(wb, corr, sd_status, v_sp):
    ws = wb.create_sheet('V_PRESUMPTION')
    header(ws, ['Corr ID', 'Basis', 'Claimant SD', 'Claimant status', 'From SP', 'To SP', 'Declared coverage',
                'From status', 'To status', 'Presumption claimable?', 'Remarks'],
           [9, 14, 12, 14, 22, 22, 14, 14, 14, 12, 40])
    n = max(len(corr), 1) + 1
    vid, vfin = f'V_SP!$A$2:$A${v_sp["n"]}', f'V_SP!${v_sp["final"]}$2:${v_sp["final"]}${v_sp["n"]}'
    for i, c in enumerate(corr, 2):
        vals = [c['corr_id'], c['basis'], c['claimant_sd'], sd_status.get(c['claimant_sd'], ''), c['from_sp'], c['to_sp'],
                c['declared_coverage']]
        for j, v in enumerate(vals, 1):
            body_cell(ws, i, j, v)
        body_cell(ws, i, 8, f'=IFERROR(INDEX({vfin},MATCH(E{i},{vid},0)),"")')
        body_cell(ws, i, 9, f'=IFERROR(INDEX({vfin},MATCH(F{i},{vid},0)),"")')
        body_cell(ws, i, 10, f'=IF(B{i}=CORRBASIS_ZA,IF(AND(D{i}=SDSTATUS_CITED,H{i}={R["PASS"]}),YN_YES,YN_NO),'
                             f'IF(H{i}={R["PASS"]},YN_YES,YN_NO))', BOLD)
        body_cell(ws, i, 11, c['remarks'])
    _rollup_colours(wb, ws, f'H2:I{n}', 'H2')
    ws.freeze_panes = 'B2'


def _v_register_view(wb, name, idx_key, nCI, keys, title):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=12)
    cols = {c.key: c for c in CI_COLS}
    header(ws, ['#'] + [cols[k].label for k in keys], [5] + [max(cols[k].width, 14) for k in keys], row=3)
    idx = _ci_rng(idx_key, nCI)
    for k in range(1, nCI):
        r = k + 3
        body_cell(ws, r, 1, k)
        for j, key in enumerate(keys, 2):
            body_cell(ws, r, j, f'=IFERROR(INDEX({_ci_rng(key, nCI)},MATCH($A{r},{idx},0))&"","")')
    ws.freeze_panes = 'B4'


def _v_checklist(wb, S, sd_id, ci_rows, v_sr):
    ws = wb.create_sheet(f'V_CL_{sd_id}'[:31])
    ws['A1'] = f'Checklist view in the clause order of {sd_id} (read-only mirror of sheet CI – enter data in CI)'
    ws['A1'].font = Font(name=F, bold=True, size=12)
    hdr = ['Clause', 'Heading', 'SR ID', 'Requirement (source)', 'SR status', 'CR ID', 'Requirement (catalogue)', 'PASS criteria',
           'Guidance', 'Object', 'Applicability', 'Evidence', 'Compliance statement', 'Assessor remarks', 'Verdict',
           'Severity', 'Reviewer remarks', 'Status', '_ci_id']
    header(ws, hdr, [14, 22, 16, 44, 14, 15, 36, 34, 30, 15, 16, 22, 30, 30, 11, 13, 26, 14, 8], row=3)
    rows = [(p['native_ref'], heading, r, cr, c) for p, heading, r, cr, c in checklist_rows(S, sd_id, ci_rows)]
    crs, guid = S.crs, S.guid
    vsr_id, vsr_st = f'V_SR!$A$2:$A${v_sr["n"]}', f'V_SR!$O$2:$O${v_sr["n"]}'
    keys = ['object_id', 'applicability', 'evidence_refs', 'client_statement', 'assessor_remarks', 'verdict',
            'finding_severity', 'reviewer_remarks', 'status']
    nCI = len(ci_rows) + 1
    cid = _ci_rng('ci_id', nCI)
    for i, (clause, heading, r, cr, c) in enumerate(rows, 4):
        body_cell(ws, i, 1, clause); body_cell(ws, i, 2, heading or None); body_cell(ws, i, 3, r['sr_id'])
        body_cell(ws, i, 4, r['statement'])
        body_cell(ws, i, 5, f'=IFERROR(INDEX({vsr_st},MATCH(C{i},{vsr_id},0)),"")')
        if cr:
            crr = crs[cr]
            g = (guid or {}).get(cr, crr['assessor_guidance'])
            for j, v in enumerate([cr, crr['statement'], crr['acceptance_criterion'], g], 6):
                body_cell(ws, i, j, v or None)
            body_cell(ws, i, 19, c['ci_id'], CHK_FONT)
            for j, k in enumerate(keys, 10):
                body_cell(ws, i, j, f'=IFERROR(INDEX({_ci_rng(k, nCI)},MATCH($S{i},{cid},0))&"","")')
        else:
            body_cell(ws, i, 6, '— no assessable item (see SR status) —', CHK_FONT)
    last = max(len(rows), 1) + 3
    reg = read_register(wb)
    add_status_colours(ws, f'R4:R{last}', '$R4', reg, ['VERDICT', 'APPL'])
    _rollup_colours(wb, ws, f'E4:E{last}', 'E4')
    ws.column_dimensions['S'].hidden = True
    ws.freeze_panes = 'D4'; ws.auto_filter.ref = f'A3:S{last}'


def _v_summary(wb, reg, prj, basis, nCI, v_cr, v_sr, v_sp):
    ws = wb.create_sheet('V_SUMMARY')
    ws.column_dimensions['A'].width = 30
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws['A1'] = f'Summary – {prj.get("Project ID", "")} – {prj.get("Product", "")}'; ws['A1'].font = Font(name=F, bold=True, size=13)
    ws['A2'] = 'All figures are live formulas. Status labels come from REGISTER.'; ws['A2'].font = Font(name=F, italic=True, size=9)
    r = 4
    statuses = [('VERDICT', e['key']) for e in reg['VERDICT']] + [('APPL', e['key']) for e in reg['APPL'] if e['key'] != 'APPLICABLE']
    ws.cell(row=r, column=1, value='Check items by object level and status').font = BOLD
    r += 1
    header(ws, ['Object level'] + [''] * len(statuses) + ['Total'], row=r)
    for j, (lst, key) in enumerate(statuses, 2):
        ws.cell(row=r, column=j, value=f'={lst}_{key}')
    hr = r
    lv, st = _ci_rng('level', nCI), _ci_rng('status', nCI)
    for e in reg['OBJLEVEL']:
        r += 1
        body_cell(ws, r, 1, f'=OBJLEVEL_{e["key"]}', BOLD)
        for j in range(2, len(statuses) + 2):
            body_cell(ws, r, j, f'=COUNTIFS({lv},$A{r},{st},{get_column_letter(j)}${hr})')
        body_cell(ws, r, len(statuses) + 2, f'=SUM(B{r}:{get_column_letter(len(statuses) + 1)}{r})', BOLD)
    r += 1
    body_cell(ws, r, 1, 'Total', BOLD)
    for j in range(2, len(statuses) + 3):
        L = get_column_letter(j); body_cell(ws, r, j, f'=SUM({L}{hr + 1}:{L}{r - 1})', BOLD)

    roll = [e['key'] for e in reg['ROLLUP']]
    def rollup_table(r, title, label_col, rng_status, rows_spec):
        ws.cell(row=r, column=1, value=title).font = BOLD
        r += 1
        header(ws, [label_col] + [''] * len(roll) + ['Total'], row=r)
        for j, k in enumerate(roll, 2):
            ws.cell(row=r, column=j, value=f'=ROLLUP_{k}')
        h = r
        for lab, extra in rows_spec:
            r += 1
            body_cell(ws, r, 1, lab, BOLD)
            for j in range(2, len(roll) + 2):
                body_cell(ws, r, j, f'=COUNTIFS({extra}{rng_status},{get_column_letter(j)}${h})')
            body_cell(ws, r, len(roll) + 2, f'=SUM(B{r}:{get_column_letter(len(roll) + 1)}{r})', BOLD)
        return r + 2

    r += 3
    r = rollup_table(r, 'Common requirements (final status)', 'Scope',
                     f'V_CR!${v_cr["final"]}$2:${v_cr["final"]}${v_cr["n"]}', [('All CRs in project', '')])
    r = rollup_table(r, 'Source requirements by source document', 'SD', f'V_SR!$O$2:$O${v_sr["n"]}',
                     [(sd, f'V_SR!$B$2:$B${v_sr["n"]},"{sd}",') for sd in basis])
    r = rollup_table(r, 'Source provisions by source document (atomic provisions)', 'SD',
                     f'V_SP!${v_sp["final"]}$2:${v_sp["final"]}${v_sp["n"]}',
                     [(sd, f'V_SP!$B$2:$B${v_sp["n"]},"{sd}",V_SP!$G$2:$G${v_sp["n"]},0,V_SP!$H$2:$H${v_sp["n"]},0,') for sd in basis])
    ws.cell(row=r, column=1, value='Findings and exclusions').font = BOLD
    r += 1
    sev = _ci_rng('finding_severity', nCI)
    for e in reg['SEV']:
        body_cell(ws, r, 1, f'=SEV_{e["key"]}', BOLD); body_cell(ws, r, 2, f'=COUNTIF({sev},A{r})'); r += 1
    body_cell(ws, r, 1, f'=APPL_NIS', BOLD); body_cell(ws, r, 2, f'=COUNTIF({_ci_rng("applicability", nCI)},A{r})')


def _snapshot(wb, C):
    for s, cols in CATALOGUE.items():
        ws = wb.create_sheet(f'cat_{s}'[:31])
        header(ws, [c.label for c in cols], [c.width for c in cols], {j: 'D9D9D9' for j in range(1, len(cols) + 1)})
        for i, rec in enumerate(C[s], 2):
            for j, c in enumerate(cols, 1):
                v = rec.get(c.key, '')
                ws.cell(row=i, column=j, value=v if v != '' else None).font = BODY
        ws.freeze_panes = 'B2'
        ws.sheet_properties.tabColor = '808080'


def _readme(wb):
    ws = wb.create_sheet('README', 0)
    ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 110
    ws['A1'] = 'Compliance Traceability – Project workbook'; ws['A1'].font = Font(name=F, size=14, bold=True)
    lines = [
        ('Where to enter data', 'PRJ, BASIS, TAGS, OBJECTS (assessor), EV (client) and the coloured input columns of CI. '
                                'Everything else is generated – re-run build_project.py after changing inputs or the catalogue.'),
        ('Column colours (CI)', 'Grey = generated · Blue = client · Yellow = assessor · Green = reviewer.'),
        ('Applicability vs verdict', 'Applicability decides IF an item is assessed (Applicable, N/A, Inherited, Not in scope (agreed)); '
                                     'verdict (Pass/Fail/Info needed/Pending) is only for Applicable items.'),
        ('Rule result (auto)', 'Applicability proposed by the CR rule from TAGS. If you change Applicability, the rebuild keeps your value.'),
        ('Evidence', 'Cite evidence IDs from sheet EV in "Evidence (IDs + section)", e.g. "P26001-EV-002 §4".'),
        ('Views', 'V_SUMMARY (counts), V_FINDINGS, V_EXCLUSIONS (for the conformity report), V_CL_<SD> (checklist in a source\'s '
                  'clause order), V_SP / V_SR / V_CR (roll-up per provision / source requirement / common requirement), '
                  'V_PRESUMPTION (Annex ZA and legal-bridge claims).'),
        ('Roll-up precedence', 'Fail > Gap > Not assessed > Open > Pass > N/A. "Inherited" counts as Pass; "Not in scope (agreed)" '
                               'gives "Not assessed"; uncovered or un-atomised provisions give "Gap".'),
        ('cat_* sheets', 'Snapshot of the catalogue baseline used for this project (read-only).'),
        ('REGISTER', 'Copy of the catalogue register. Change category strings in the catalogue, then rebuild.'),
    ]
    for i, (a, b) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=a).font = BOLD
        c = ws.cell(row=i, column=2, value=b); c.font = BODY; c.alignment = WRAP


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('catalogue'); ap.add_argument('project')
    ap.add_argument('--init', action='store_true')
    ap.add_argument('--seed', help='project seed JSON (e.g. projects/P26901.json) to pre-fill the input sheets on --init')
    a = ap.parse_args()
    if a.init:
        print(init_project(a.catalogue, a.project, load_project_seed(a.seed) if a.seed else None))
    else:
        print(build_project(a.catalogue, a.project))
