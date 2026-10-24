"""ctm.project on the JSON store must give the same selection as build_project.py on the workbooks (pytest)."""
import os, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import pytest                                                                         # noqa: E402
from openpyxl import Workbook, load_workbook                                          # noqa: E402

import build_project, make_catalogue                                                  # noqa: E402
from ctm.common import read_table                                                     # noqa: E402
from ctm.project import checklist_rows, register_from_vocab, select_project           # noqa: E402
from ctm.register import read_register, write_register                                # noqa: E402
from ctm.schema import CI_COLS                                                        # noqa: E402
from ctm.store import load_catalogue, load_project_seed, seed_inputs, to_tables       # noqa: E402

FIX = os.path.join(ROOT, 'tests', 'fixtures')
SAMPLE = os.path.join(FIX, 'sample', 'catalogue')
SEED = os.path.join(FIX, 'sample_project.json')


@pytest.fixture(scope='module')
def built(tmp_path_factory):
    d = tmp_path_factory.mktemp('prj')
    cat, prj = str(d / 'cat.xlsx'), str(d / 'prj.xlsx')
    tables = to_tables(load_catalogue(SAMPLE, []))
    make_catalogue.build(cat, tables)
    build_project.init_project(cat, prj, load_project_seed(SEED))
    build_project.build_project(cat, prj)
    _, basis, tags, objects, _ = seed_inputs(load_project_seed(SEED))
    sel = select_project(tables, register_from_vocab(), [b['sd_id'] for b in basis], tags, objects)
    return load_workbook(prj), sel


def test_register_from_vocab_matches_workbook_register():
    wb = Workbook(); wb.remove(wb.active)
    write_register(wb)
    assert register_from_vocab() == read_register(wb)


def test_cis_match_workbook(built):
    wb, sel = built
    ws_rows = {r['ci_id']: r for r in read_table(wb, 'CI', CI_COLS)}
    assert [c['ci_id'] for c in sel.ci_rows] == list(ws_rows)
    for c in sel.ci_rows:
        w = ws_rows[c['ci_id']]
        assert (c['rule_applicability'], c['applicability_justification']) == \
               (w['rule_applicability'], w['applicability_justification'] or '')
    assert sel.subst                                # the fixture exercises full-refinement substitution


def test_checklist_rows_match_views(built):
    wb, sel = built
    for sd in sel.basis:
        ws = wb[f'V_CL_{sd}'[:31]]
        sheet = [(r[2], r[5] if r[18] else None, r[18]) for r in ws.iter_rows(min_row=4, values_only=True) if r[2]]
        rows = [(r['sr_id'], cr, c['ci_id'] if c else None) for _, _, r, cr, c in checklist_rows(sel, sd)]
        assert rows == sheet, sd
